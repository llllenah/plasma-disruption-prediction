"""Презентація для захисту (PPTX, 16:9, темна тема, Arial — однаково на Mac і Windows).

    python work/make_deck.py presentation/plasma_disruption.pptx
"""
import sys
from pathlib import Path

import qrcode
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "presentation/plasma_disruption.pptx")
OUT.parent.mkdir(parents=True, exist_ok=True)
REPO = "github.com/llllenah/plasma-disruption-prediction"
FONT = "Arial"

BG, CARD, CARD2 = "0B1120", "16203A", "1F2B4A"
TEXT, MUTED, DIM = "F1F5F9", "94A3B8", "64748B"
VIOLET, INDIGO, EMERALD, ORANGE, RED, AMBER, CYAN = "8B5CF6", "6366F1", "10B981", "F97316", "EF4444", "F59E0B", "22D3EE"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 15
_n = [0]


def rgb(h):
    return RGBColor.from_string(h)


def tb(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margin=0.0, name=None):
    """paras: [(текст | [(текст, {опції})], {size, color, bold, align, after, italic})]."""
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    for i, (content, o) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(o.get("after", 0))
        runs = content if isinstance(content, list) else [(content, {})]
        for text, ro in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = FONT
            f.size = Pt(ro.get("size", o.get("size", 16)))
            f.bold = ro.get("bold", o.get("bold", False))
            f.italic = ro.get("italic", o.get("italic", False))
            f.color.rgb = rgb(ro.get("color", o.get("color", TEXT)))
    return s


def box(slide, x, y, w, h, fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, line=None, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.5)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def boxed_text(slide, x, y, w, h, paras, fill=CARD, pad=0.22, anchor=MSO_ANCHOR.TOP, line=None):
    box(slide, x, y, w, h, fill, line=line)
    return tb(slide, x + pad, y + pad * 0.8, w - 2 * pad, h - 1.6 * pad, paras, anchor=anchor)


def badge(slide, x, y, d, text, fill, size=16, color=BG):
    s = box(slide, x, y, d, d, fill, shape=MSO_SHAPE.OVAL)
    tb(slide, x, y, d, d, [(text, dict(size=size, bold=True, color=color, align=PP_ALIGN.CENTER))], anchor=MSO_ANCHOR.MIDDLE)
    return s


def arrow(slide, x, y, w=0.45, h=0.32, fill=DIM):
    box(slide, x, y, w, h, fill, shape=MSO_SHAPE.RIGHT_ARROW)


def down_arrow(slide, x, y, w=0.32, h=0.4, fill=DIM):
    box(slide, x, y, w, h, fill, shape=MSO_SHAPE.DOWN_ARROW)


def new_slide(tag, title, tag_color=VIOLET, notes=""):
    _n[0] += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    if tag:
        tb(s, 0.6, 0.38, 9.0, 0.32, [(tag, dict(size=12, bold=True, color=tag_color))], name="tag")
    if title:
        tb(s, 0.6, 0.66, 12.1, 0.9, [(title, dict(size=32, bold=True))], name="title")
    tb(s, 12.2, 7.0, 0.6, 0.3, [(f"{_n[0]}/{TOTAL}", dict(size=10, color=DIM, align=PP_ALIGN.RIGHT))])
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def no_fill_chart(chart):
    cs = chart._chartSpace
    for parent in (cs, cs.find(qn("c:chart")).find(qn("c:plotArea"))):
        sp = parent.find(qn("c:spPr"))
        if sp is None:
            sp = etree.SubElement(parent, qn("c:spPr"))
            if parent is cs:
                cs.find(qn("c:chart")).addnext(sp)
        for ch in list(sp):
            sp.remove(ch)
        etree.SubElement(sp, qn("a:noFill"))
        ln = etree.SubElement(sp, qn("a:ln"))
        etree.SubElement(ln, qn("a:noFill"))


def bar_chart(slide, x, y, w, h, cats, vals, colors, fmt="0.000", vmin=None, vmax=None, horizontal=False, label_size=14):
    data = CategoryChartData()
    data.categories = cats
    data.add_series("v", vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name = FONT
    ch.font.color.rgb = rgb(MUTED)
    plot = ch.plots[0]
    plot.gap_width = 55
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold, dl.font.color.rgb = Pt(label_size), True, rgb(TEXT)
    for i, c in enumerate(colors):
        pt = plot.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
    va, ca = ch.value_axis, ch.category_axis
    if vmin is not None:
        va.minimum_scale = vmin
    if vmax is not None:
        va.maximum_scale = vmax
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("26334F")
    va.visible = False
    ca.tick_labels.font.size, ca.tick_labels.font.color.rgb = Pt(12), rgb(MUTED)
    ca.format.line.color.rgb = rgb("334155")
    if horizontal:
        ca.reverse_order = True
    no_fill_chart(ch)
    return ch


def table(slide, x, y, w, rows, col_w, row_h=0.42, header_fill=CARD2, fills=None, sizes=None, colors=None, bold_first_col=True):
    nr, nc = len(rows), len(rows[0])
    gs = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gs.table
    tblPr = gs._element.graphic.graphicData.tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            f = header_fill if i == 0 else (fills[i][j] if fills else CARD)
            c.fill.fore_color.rgb = rgb(f)
            c.margin_left = c.margin_right = Inches(0.1)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            r.font.name = FONT
            r.font.size = Pt(sizes[i][j] if sizes else (13 if i == 0 else 14))
            r.font.bold = i == 0 or (j == 0 and bold_first_col)
            r.font.color.rgb = rgb(colors[i][j] if colors else (MUTED if i == 0 else TEXT))
    return gs


def chips(slide, x, y, items, h=0.42, gap=0.18, size=13):
    cx = x
    for text, color in items:
        w = 0.32 + 0.105 * len(text)
        box(slide, cx, y, w, h, CARD2, radius=0.5)
        tb(slide, cx, y, w, h, [(text, dict(size=size, bold=True, color=color, align=PP_ALIGN.CENTER))], anchor=MSO_ANCHOR.MIDDLE)
        cx += w + gap


# ───────────────────────── 1. титул ─────────────────────────
s = new_slide("ХАКАТОН ФМФ КПІ · 9–11 ЖОВТНЯ 2026", None, notes=(
    "≈30 с. Привітання. Ми розв'язували задачу організаторів: за 100 приладами токамака в реальному часі передбачити "
    "зрив плазми. Зробили три речі: модель тривоги, нейромережу, яка пояснює причину, і симулятор, де модель сама "
    "змінює параметри, щоб зриву не було."))
for i, (d, c) in enumerate(((3.9, VIOLET), (3.1, ORANGE), (2.3, INDIGO))):
    o = box(s, 10.85 - d / 2, 2.85 - d * 0.36, d, d * 0.72, None, shape=MSO_SHAPE.OVAL, line=c)
    o.line.width = Pt(3 - i * 0.6)
tb(s, 0.6, 1.5, 8.2, 2.4, [("Передбачаємо, пояснюємо і запобігаємо зриву плазми", dict(size=44, bold=True))])
tb(s, 0.6, 3.75, 7.8, 0.9, [("Онлайн-модель тривоги · нейромережа причин ImpactNet · симулятор з автокеруванням",
                             dict(size=18, color=MUTED))])
stats = (("0,924", "об'єм заліку\n(бейзлайн 0,765)", EMERALD), ("49 / 49", "причин зривів\nвгадано", VIOLET),
         ("11 → 5", "зривів з\nавтокеруванням", ORANGE))
for i, (big, small, c) in enumerate(stats):
    x = 0.6 + i * 2.75
    boxed_text(s, x, 4.85, 2.5, 1.45, [(big, dict(size=30, bold=True, color=c, after=2)),
                                       (small, dict(size=12, color=MUTED))])
tb(s, 0.6, 6.55, 12, 0.5, [("Дмитрій Почко · Олександр Чорій · Ілля Симоненко · Олена Сергієнко · Софія Щербина · Арина Яковець",
                            dict(size=13, color=MUTED))])

# ───────────────────────── 2. проблема ─────────────────────────
s = new_slide("ПРОБЛЕМА", "Зрив — найдорожча аварія токамака", notes=(
    "≈45 с. Зрив — плазма за мілісекунди втрачає утримання, енергія б'є в стінку й котушки. У симуляторі це "
    "пуассонівський кубик: що ближче до операційної межі, то вища інтенсивність λ. Після кидка 50–100 мс видно "
    "передвісник, потім гасіння. Модель бачить 100 приладів кожні 5 мс, лише минуле. Тривога зараховується, якщо прийшла "
    "щонайменше за lead до гасіння, а хибних тривог не більше бюджету. Головна метрика — об'єм: середня якість по всіх "
    "запасах і бюджетах."))
y0 = 2.35
box(s, 0.7, y0 + 1.05, 7.4, 0.07, "334155", shape=MSO_SHAPE.RECTANGLE)
phases = ((0.7, 3.2, "плазма наближається\nдо межі: λ росте", AMBER), (3.95, 1.5, "кидок\nкубика", RED),
          (5.5, 1.6, "передвісник\n50–100 мс", ORANGE), (7.15, 0.95, "гасіння\n1–2 мс", RED))
for x, w, label, c in phases:
    box(s, x, y0 + 0.88, w, 0.42, c, radius=0.5)
    tb(s, x - 0.1, y0 + 1.45, w + 0.2, 0.8, [(label, dict(size=13, color=TEXT, align=PP_ALIGN.CENTER))])
badge(s, 4.55, y0 - 0.05, 0.62, "!", EMERALD, size=20)
tb(s, 2.6, y0 - 0.65, 4.6, 0.5, [("тривога моделі — щонайменше за lead до гасіння", dict(size=13, color=EMERALD, align=PP_ALIGN.CENTER))])
for i, (h, t, c) in enumerate((("100 приладів, крок 5 мс", "магнітні, інтерферометр, ECE, болометри, SXR, газ, нагрів", CYAN),
                                ("Лише минуле", "модель працює онлайн, у реальному часі", VIOLET),
                                ("Тривога ≥ lead до гасіння", "хибних тривог — не більше бюджету", EMERALD),
                                ("Метрика — об'єм", "середня якість по всіх запасах і бюджетах", ORANGE))):
    yy = 1.75 + i * 1.2
    box(s, 8.75, yy, 4.0, 1.02, CARD)
    badge(s, 8.95, yy + 0.27, 0.48, str(i + 1), c, size=15)
    tb(s, 9.6, yy + 0.12, 3.05, 0.85, [(h, dict(size=15, bold=True, after=2)), (t, dict(size=11, color=MUTED))])
tb(s, 0.7, 5.3, 7.4, 1.4, [("Задача: побачити зрив якомога раніше, але не «лякати» тихі розряди хибними тривогами.",
                            dict(size=18, color=TEXT, bold=True))])

# ───────────────────────── 3. бейзлайн ─────────────────────────
s = new_slide("АНАЛІЗ БЕЙЗЛАЙНУ", "Бейзлайн ловить пізно і спирається на одну величину", notes=(
    "≈40 с. Бейзлайн організаторів — поріг на частці межі Грінвальда плюс сигнал locked mode. Об'єм 0,765. При 30 мс і 3 % "
    "хибних він ловить лише 59 % зривів. Причини: одна фізична величина, немає пам'яті, шум приладів б'є прямо в поріг."))
boxed_text(s, 0.6, 1.8, 4.6, 4.6, [("0,765", dict(size=72, bold=True, color=ORANGE)),
                                   ("об'єм заліку бейзлайну", dict(size=15, color=MUTED, after=18)),
                                   ("59 %", dict(size=44, bold=True, color=ORANGE)),
                                   ("зривів спіймано за 30 мс\nпри 3 % хибних тривог", dict(size=15, color=MUTED))], pad=0.35)
for i, (h, t, c) in enumerate((("Що робить", "поріг на f_GW (частка межі густини) + сигнал locked mode із сідлової петлі", CYAN),
                                ("Чому слабко", "одна величина, без пам'яті про розряд; шум і дрейф приладів напряму смикають поріг; зриви від q95 і радіації майже не бачить", RED),
                                ("Наша ціль", "ловити раніше й більше, не піднімаючи хибні тривоги — і пояснити, чому плазма зривається", EMERALD))):
    yy = 1.8 + i * 1.55
    box(s, 5.55, yy, 7.2, 1.38, CARD)
    badge(s, 5.8, yy + 0.42, 0.55, "→" if i == 2 else ("?" if i == 1 else "i"), c, size=17)
    tb(s, 6.55, yy + 0.18, 6.0, 1.1, [(h, dict(size=17, bold=True, after=3)), (t, dict(size=13, color=MUTED))])

# ───────────────────────── 4. гіпотези ─────────────────────────
s = new_slide("ГІПОТЕЗИ", "Гіпотези: чотири «так», одне важливе «ні»", notes=(
    "≈50 с. H1: два детектори — пізній упевнений передвісник і рання тривога. H2: формули зі шпаргалки не потрібні. "
    "H3: мітка сценарію в метаданих — підказка симулятора, а не фізика плазми. H4 спростована: навіть точна λ — погана "
    "сигналізація, бо зрив — кидок кубика; звідси дві моделі. H5: фізичне априорі дає розпізнати причини без розмітки."))
hyps = (("H1", "Два детектори краще за один", "пізній упевнений передвісник + рання тривога", True),
        ("H2", "Формули шпаргалки не потрібні", "сирих згладжених каналів досить: 0,924 без них", True),
        ("H3", "Мітка сценарію — підказка, не фізика", "прибрали — об'єм упав лише на ~0,01", True),
        ("H4", "Знаючи точну λ, можна бити тривогу", "ні: навіть справжня λ дає 0,78 — зрив є кубик", False),
        ("H5", "Фізичне априорі робить причини видимими", "без розмітки: 49/49 з априорі, 41/49 без", True))
for i, (k, h, t, ok) in enumerate(hyps):
    col, row = (i % 3, i // 3)
    x = 0.6 + col * 4.12 + (2.06 if row == 1 else 0)
    yy = 1.8 + row * 2.45
    box(s, x, yy, 3.9, 2.2, CARD)
    badge(s, x + 0.22, yy + 0.22, 0.6, k, VIOLET if ok else ORANGE, size=14)
    box(s, x + 2.55, yy + 0.3, 1.15, 0.42, EMERALD if ok else ORANGE, radius=0.5)
    tb(s, x + 2.55, yy + 0.3, 1.15, 0.42, [("так" if ok else "ні", dict(size=13, bold=True, color=BG, align=PP_ALIGN.CENTER))],
       anchor=MSO_ANCHOR.MIDDLE)
    tb(s, x + 0.22, yy + 0.98, 3.5, 1.15, [(h, dict(size=15, bold=True, after=4)), (t, dict(size=12, color=MUTED))])

# ───────────────────────── 5. модель тривоги ─────────────────────────
s = new_slide("РІШЕННЯ 1 · МОДЕЛЬ ТРИВОГИ", "Дві голови бустингу: передвісник і рання тривога", EMERALD, notes=(
    "≈45 с. Ознаки рахуються онлайн: згладжування на трьох масштабах, тренди, максимуми, детектор передвісника на сідловій "
    "петлі й котушках Мірнова, команди актуаторів. Голова A вчиться на «зрив ≤ 0,1 с» — пізно, але впевнено. Голова B — "
    "«зрив ≤ 4,5 с» — рання тривога. Беремо максимум, на запланованому спаді струму нуль, далі згладжування. Дерева "
    "перекладені в масиви numpy: крок 0,9 мс на ноутбуці."))
flow = ((0.6, "100 приладів", "кожні 5 мс", CYAN), (3.32, "Онлайн-ознаки", "EMA 50 мс / 250 мс / 1 с,\nтренди, передвісник,\nактуатори", VIOLET))
for x, h, t, c in flow:
    box(s, x, 2.4, 2.15, 2.2, CARD, line=c)
    tb(s, x + 0.12, 2.5, 1.91, 2.0, [(h, dict(size=16, bold=True, color=c, after=6, align=PP_ALIGN.CENTER)),
                                     (t, dict(size=12, color=MUTED, align=PP_ALIGN.CENTER))], anchor=MSO_ANCHOR.MIDDLE)
arrow(s, 2.8, 3.35)
arrow(s, 5.52, 3.35, w=0.4)
for yy, h, t in ((1.85, "Голова A", "«зрив ≤ 0,1 с»\nпередвісник"), (3.85, "Голова B", "«зрив ≤ 4,5 с»\nрання тривога")):
    box(s, 6.0, yy, 2.5, 1.55, CARD, line=EMERALD)
    tb(s, 6.1, yy + 0.08, 2.3, 1.4, [(h, dict(size=16, bold=True, color=EMERALD, after=4, align=PP_ALIGN.CENTER)),
                                     (t, dict(size=12, color=MUTED, align=PP_ALIGN.CENTER))], anchor=MSO_ANCHOR.MIDDLE)
arrow(s, 8.65, 3.35)
box(s, 9.2, 2.4, 3.55, 2.2, CARD2, line=ORANGE)
tb(s, 9.35, 2.5, 3.25, 2.0, [("p = max(A, 0,3·B)", dict(size=17, bold=True, color=ORANGE, after=6, align=PP_ALIGN.CENTER)),
                             ("0 на запланованому спаді струму\n→ згладжування → бігучий максимум\n→ тривога, коли p > поріг",
                              dict(size=12, color=MUTED, align=PP_ALIGN.CENTER))], anchor=MSO_ANCHOR.MIDDLE)
chips(s, 0.6, 5.55, (("лише numpy", CYAN), ("крок 0,9 мс", EMERALD), ("4 500 розрядів", VIOLET),
                     ("без формул шпаргалки", ORANGE), ("без мітки сценарію", ORANGE)))
tb(s, 0.6, 6.2, 12, 0.6, [("Поріг підбирає сам залік під заданий бюджет хибних тривог — модель не треба перенавчати під налаштування журі.",
                           dict(size=13, color=MUTED))])

# ───────────────────────── 6. результат ─────────────────────────
s = new_slide("РЕЗУЛЬТАТ ПРОТИ БЕЙЗЛАЙНУ", "0,765 → 0,924 без формул і 98 % зривів за 30 мс", EMERALD, notes=(
    "≈50 с. Офіційний залік toksym на 4 наборах, 747 розрядів, 228 зривів. Бейзлайн 0,765, наша чиста модель 0,914 на 1800 "
    "розрядах і 0,924 на 4500. Для довідки модель з фізикою — 0,944, але вона використовує формули шпаргалки. При 30 мс і 5 % "
    "хибних ловимо 98 %. Приріст від даних перевірено парним бутстрепом: +0,010, P = 0,97."))
bar_chart(s, 0.5, 1.75, 7.6, 4.9, ["Бейзлайн", "Наша\n(1 800 розр.)", "Наша фінальна\n(4 500 розр.)", "З фізикою\n(довідково)"],
          [0.765, 0.914, 0.924, 0.944], [ORANGE, INDIGO, EMERALD, DIM], vmin=0.6, vmax=1.0)
for i, (big, small, c) in enumerate((("98,2 %", "зривів за 30 мс\nпри 5 % хибних", EMERALD), ("+0,159", "об'єму до бейзлайну", VIOLET),
                                      ("0,9 мс", "крок на ноутбуці\n(бюджет 5 мс)", CYAN))):
    boxed_text(s, 8.6, 1.75 + i * 1.65, 4.15, 1.45, [(big, dict(size=34, bold=True, color=c, after=2)),
                                                    (small, dict(size=12, color=MUTED))])
tb(s, 0.6, 6.7, 12, 0.4, [("Офіційний toksym score: сиди 1, 1 shift, 20261001, 42 · 747 розрядів · 228 зривів", dict(size=11, color=DIM))])

# ───────────────────────── 7. стійкість ─────────────────────────
s = new_slide("СТІЙКІСТЬ", "Стійка до будь-яких налаштувань журі", EMERALD, notes=(
    "≈35 с. Запас і бюджет вибирає команда, журі може поставити будь-які. Таблиця — частка зривів, спійманих вчасно, на "
    "тих самих 4 наборах. До 30 мс — 98–100 % за будь-якого бюджету. Довгі запаси при жорсткому бюджеті — природна межа: "
    "зрив — кубик. Для уникнення зриву (0,3 с) при 10–15 % хибних ловимо 67–82 %."))
grid = [["запас \\ хибні тривоги", "3 %", "5 %", "10 %", "15 %"],
        ["10 мс", 99.1, 99.1, 99.6, 99.6], ["30 мс", 97.8, 98.2, 99.6, 99.6], ["0,1 с", 39.9, 50.9, 71.9, 83.3],
        ["0,3 с", 36.0, 47.8, 66.7, 81.6], ["1 с", 22.4, 29.8, 47.4, 61.0]]


def heat(v):
    return "065F46" if v >= 95 else "047857" if v >= 80 else "15803D" if v >= 65 else "3F6212" if v >= 50 else "713F12" if v >= 35 else "7C2D12"


fills = [[None] * 5] + [[CARD] + [heat(v) for v in r[1:]] for r in grid[1:]]
rows = [grid[0]] + [[r[0]] + [f"{v:.1f} %".replace(".", ",") for v in r[1:]] for r in grid[1:]]
table(s, 0.9, 1.85, 8.4, rows, [2.4, 1.5, 1.5, 1.5, 1.5], row_h=0.72, fills=fills,
      sizes=[[14] * 5] + [[16] * 5] * 5)
for i, (h, t, c) in enumerate((("Пом'якшення зриву", "запас 10–30 мс, 5 % хибних:\n98–99 % вчасно", EMERALD),
                                ("Уникнення зриву", "запас 0,3 с, 10–15 % хибних:\n67–82 % вчасно", AMBER))):
    boxed_text(s, 9.75, 1.85 + i * 1.9, 3.0, 1.65, [(h, dict(size=16, bold=True, color=c, after=6)), (t, dict(size=13, color=MUTED))])
tb(s, 0.9, 6.35, 12, 0.5, [("Частка зривів, спійманих вчасно · 4 офіційні набори · перерахунок звітів toksym rescore", dict(size=11, color=DIM))])

# ───────────────────────── 8. чесність ─────────────────────────
s = new_slide("ЧЕСНА МОДЕЛЬ", "Прибрали всі підказки — якість майже не впала", EMERALD, notes=(
    "≈45 с. Тімейт помітив: мітка сценарію в метаданих — scen_none гарантує, що зриву не буде. Це підказка симулятора, "
    "а не плазма. Ми прибрали і її, і таймери програми, і всі формули шпаргалки. Ціна — близько сотої об'єму. "
    "Перевірка чесна: ознаки лише з минулого, порівняння на тих самих 747 розрядах парним бутстрепом, правило P > 0,9 "
    "зафіксоване наперед, ноутбук і сервер дають однакові цифри."))
rows = [["що прибрали", "чому", "ціна"], ["Формули шпаргалки (f_GW, q95, βN, P_rad/P_in, λ)", "правило задачі", "0,944 → 0,924"],
        ["Мітка сценарію з meta", "scen_none ⇒ зриву немає: підказка", "≈ −0,01"],
        ["Таймери від подій програми", "розклад симулятора", "≈ 0"]]
table(s, 0.6, 1.85, 7.7, rows, [3.3, 2.5, 1.9], row_h=0.78, sizes=[[13] * 3] + [[13, 13, 15]] * 3)
for i, (h, t, c) in enumerate((("Лише минуле", "той самий клас ознак у навчанні і в заліку", CYAN),
                                ("Парний бутстреп", "ті самі 747 розрядів; правило P > 0,9 наперед", VIOLET),
                                ("Детермінованість", "ноутбук і сервер: однакові результати", EMERALD))):
    yy = 1.85 + i * 1.5
    box(s, 8.75, yy, 4.0, 1.3, CARD)
    badge(s, 8.95, yy + 0.38, 0.52, "✓", c, size=16)
    tb(s, 9.65, yy + 0.18, 3.0, 1.05, [(h, dict(size=15, bold=True, after=3)), (t, dict(size=12, color=MUTED))])
tb(s, 0.6, 5.35, 7.7, 1.2, [("Модель вчить плазму, а не розклад симулятора: якість тримається і без підказок.",
                             dict(size=18, bold=True))])

# ───────────────────────── 9. фізика ─────────────────────────
s = new_slide("ФІЗИКА ЗРИВУ", "Зрив — кубик від чотирьох операційних меж", AMBER, notes=(
    "≈50 с. Фізика симулятора: інтенсивність зриву λ = 1 − добуток (1 − rᵢ³), де rᵢ — сигмоїди відстані до чотирьох "
    "меж: Грінвальда за густиною, кінку за q95, радіаційного колапсу і Тройона за βN. Кожну межу можна відсунути "
    "актуатором, але дії конфліктують: струм тягне в обидва боки — більший Iₚ віддаляє від Грінвальда, але наближає "
    "кінк. Неон накопичується пів секунди, густина реагує за 0,3 с — діяти треба завчасно."))
boxed_text(s, 0.6, 1.75, 12.15, 0.8, [([("λ = 1 − Π(1 − rᵢ³)·(1 − r_lm),   ", dict(size=20, bold=True, color=AMBER)),
                                        ("rᵢ = σ((xᵢ − cᵢ)/wᵢ)  —  м'який поріг: далеко від межі ≈ 0, на межі 0,5", dict(size=15, color=MUTED))],
                                       {})], anchor=MSO_ANCHOR.MIDDLE, pad=0.25)
rows = [["механізм", "драйвер", "межа", "що зменшує ризик", "ціна"],
        ["Грінвальд", "f_GW = n̄ₑ / n_GW", "f_GW ≈ 1", "↓ уставка густини, ↓ D₂; ↑ Iₚ", "↑ Iₚ наближає кінк"],
        ["Кінк", "q95", "q95 ≈ 2,44", "↓ Iₚ; ↑ Bₜ, κ, δ", "↓ Iₚ наближає Грінвальда"],
        ["Радіація", "P_rad / P_in", "≈ 0,8", "неон → 0; ↑ нагрів; ↓ густина", "нагрів піднімає βN"],
        ["Тройон", "βN", "βN ≈ 2,4", "↓ нагрів; ↑ Iₚ, Bₜ", "менше утримання"]]
table(s, 0.6, 2.8, 12.15, rows, [1.8, 2.2, 1.6, 3.6, 2.95], row_h=0.62, sizes=[[13] * 5] + [[14] * 5] * 4)
tb(s, 0.6, 6.05, 12.1, 0.9, [([("Iₚ тягне в обидва боки", dict(bold=True, color=AMBER)),
                               (" · густина реагує за ~0,3 с · неон накопичується за ~0,5 с — керувати треба завчасно, а не в момент тривоги",
                                dict(color=MUTED))], dict(size=14))])

# ───────────────────────── 10. фізичні ознаки проти динаміки ─────────────────────────
s = new_slide("ФІЗИЧНІ ОЗНАКИ ПРОТИ ДИНАМІКИ", "Пам'ять у часі важить більше за фізичні проксі", AMBER, notes=(
    "≈45 с. Ми перевірили альтернативи на XGBoost: сирі канали, плюс фізичні агрегати, плюс динаміка — лаги, різниці, "
    "ковзні статистики, і варіант із формулами-проксі. На тих самих 200 розрядах інтерфейсу при запасі 0,1 с фізичні "
    "проксі не додають нічого: 25 %. Динаміка дає 35 %. Наша модель з пам'яттю на трьох масштабах і детектором "
    "передвісника — 63 %. Висновок: головне — історія сигналів, а не формули."))
bar_chart(s, 0.5, 1.75, 7.8, 4.7, ["XGBoost\nфізичні ознаки", "XGBoost\nформули-проксі", "XGBoost\n+ динаміка", "Наша модель"],
          [25.4, 25.4, 34.9, 63.5], [DIM, DIM, AMBER, EMERALD], fmt='0.0" %"', vmin=0, vmax=75)
for i, (h, t, c) in enumerate((("Фізичні проксі ≈ 0", "агрегати й формули майже нічого не додають до сирих каналів", DIM),
                                ("Динаміка — головне", "лаги, тренди, ковзні статистики: +10 п.п.", AMBER),
                                ("Наш дизайн", "EMA на 3 масштабах + детектор передвісника: 63 %", EMERALD))):
    boxed_text(s, 8.65, 1.75 + i * 1.55, 4.1, 1.38, [(h, dict(size=15, bold=True, color=c, after=4)), (t, dict(size=12, color=MUTED))])
tb(s, 0.6, 6.55, 12, 0.4, [("Ті самі 200 розрядів інтерфейсу · запас 0,1 с · 5 % хибних тривог · частка зривів, спійманих вчасно",
                            dict(size=11, color=DIM))])

# ───────────────────────── 11. ImpactNet ─────────────────────────
s = new_slide("РІШЕННЯ 2 · IMPACTNET", "ImpactNet називає причину кожного зриву", VIOLET, notes=(
    "≈55 с. Друга модель відповідає не «коли», а «чому». Рівно 5 вихідних нейронів — Грінвальд, q95, радіація, Тройон, "
    "інші. Кожен нейрон — формула ризику симулятора плюс поправка мережі. Тому частки складаються рівно в 100 %, а "
    "загальний ризик — та сама λ, що в симуляторі. Всередині згортки бачать 155 мс, GRU дає пам'ять розряду. "
    "Мітки — справжній ризик зі стану плазми, але на вхід йдуть лише 100 приладів. Результат: причину вгадано в 49 з 49 "
    "зривів, R² загального ризику 0,996, крок 0,6 мс на чистому numpy."))
steps = (("218 входів", "100 приладів\n+ маски відмов\n+ 18 фіз. величин", CYAN), ("Згортки", "причинні,\nвікно 155 мс", INDIGO),
         ("GRU (128)", "пам'ять\nрозряду", VIOLET), ("5 нейронів", "+ фізичне\nаприорі", ORANGE))
for i, (h, t, c) in enumerate(steps):
    x = 0.6 + i * 2.0
    box(s, x, 1.85, 1.7, 1.95, CARD, line=c)
    tb(s, x + 0.08, 1.9, 1.54, 1.85, [(h, dict(size=14, bold=True, color=c, after=5, align=PP_ALIGN.CENTER)),
                                      (t, dict(size=11, color=MUTED, align=PP_ALIGN.CENTER))], anchor=MSO_ANCHOR.MIDDLE)
    if i < 3:
        arrow(s, x + 1.73, 2.65, w=0.24, h=0.28)
mech = (("Грінвальд", ORANGE), ("q95", INDIGO), ("радіація", AMBER), ("Тройон", CYAN), ("інші", DIM))
for i, (m, c) in enumerate(mech):
    box(s, 0.6 + i * 1.57, 4.1, 1.45, 0.45, c, radius=0.5)
    tb(s, 0.6 + i * 1.57, 4.1, 1.45, 0.45, [(m, dict(size=13, bold=True, color=BG, align=PP_ALIGN.CENTER))], anchor=MSO_ANCHOR.MIDDLE)
boxed_text(s, 0.6, 4.8, 7.7, 1.6, [("частки = 100·Cₖ / ΣC,  Cₖ = −ln(1 − rₖᵖ)  →  сума рівно 100 %", dict(size=14, color=TEXT, after=4)),
                                   ("ризик λ = 1 − e^(−ΣC)  —  та сама формула, що в симуляторі", dict(size=14, color=TEXT, after=4)),
                                   ("P(зрив за h с) = 1 − e^(−λh)", dict(size=14, color=TEXT))])
for i, (big, small, c) in enumerate((("49 / 49", "головна причина\nв момент зриву", VIOLET), ("0,996", "R² загального ризику λ", EMERALD),
                                      ("0,6 мс", "крок, лише numpy", CYAN))):
    boxed_text(s, 8.75, 1.85 + i * 1.55, 4.0, 1.38, [(big, dict(size=32, bold=True, color=c, after=2)), (small, dict(size=12, color=MUTED))])
tb(s, 0.6, 6.6, 12, 0.4, [("200 тестових розрядів, яких мережа не бачила · 20 зривів від Грінвальда, 15 від q95, 14 від радіації — усі вгадані",
                           dict(size=11, color=DIM))])

# ───────────────────────── 12. інсайт ─────────────────────────
s = new_slide("ГОЛОВНИЙ ІНСАЙТ", "Точна λ не попереджає: потрібні дві моделі", VIOLET, notes=(
    "≈50 с. Якщо використати ImpactNet як будильник, об'єм лише 0,77. Але й справжня λ симулятора дає 0,78 — мережа "
    "відтворює її майже ідеально. Причина: зрив — кидок кубика, тихі розряди теж підходять до межі. Тому «коли» — модель "
    "тривоги, яка бачить передвісник після кидка, а «чому» — ImpactNet. І друге: на справжньому реакторі λ ніхто не знає. "
    "Ми навчили ту саму мережу лише з моментів зривів — і вона знову вгадала 49 з 49 причин, але тільки з фізичним "
    "априорі. Без нього — 41."))
bar_chart(s, 0.5, 1.75, 6.3, 4.6, ["Модель\nтривоги", "ImpactNet\nяк будильник", "Справжня λ\n(оракул)"],
          [0.920, 0.767, 0.780], [EMERALD, VIOLET, DIM], vmin=0.6, vmax=1.0)
tb(s, 0.6, 6.4, 6.2, 0.6, [("Об'єм як сигналізація, ті самі 200 тестових розрядів", dict(size=11, color=DIM))])
boxed_text(s, 7.1, 1.75, 5.65, 1.75, [("Дві ролі", dict(size=16, bold=True, color=VIOLET, after=5)),
                                      ("«коли» — модель тривоги: бачить передвісник після кидка кубика\n«чому» — ImpactNet: розкладає ризик на механізми",
                                       dict(size=13, color=MUTED))])
rows = [["навчання без розмітки λ", "причин\nвгадано", "R² λ"], ["з мітками λ (основна)", "49 / 49", "0,996"],
        ["лише моменти зривів + априорі", "49 / 49", "0,948"], ["без фізичного априорі", "41 / 49", "0,660"]]
table(s, 7.1, 3.75, 5.65, rows, [3.15, 1.3, 1.2], row_h=0.6, sizes=[[12] * 3] + [[13, 15, 15]] * 3)
tb(s, 7.1, 6.3, 5.65, 0.7, [("Як на справжньому реакторі: λ невідома, є лише моменти зривів", dict(size=12, color=MUTED))])

# ───────────────────────── 13. симулятор ─────────────────────────
s = new_slide("СИМУЛЯТОР З КЕРУВАННЯМ", "Модель сама керує розрядом у реальному часі", ORANGE, notes=(
    "≈45 с. Ми зробили симулятор, де модель під'єднана до розряду. Панель показує ймовірність зриву і живі частки "
    "механізмів. Порадник бере головну загрозу і пропонує конкретну дію — наприклад, знизити уставку густини до 0,62 і "
    "D₂ удвічі. Автокерування виконує дію, коли p вище порогу, з паузою, бо плазмі треба час відреагувати. Плюс 10 "
    "повзунків для ручного керування."))
panel = ((0.6, 1.85, 3.9, 1.6, "Прогноз зриву", "p = 0,41 · λ = 0,52\nграфік p у реальному часі", ORANGE),
         (0.6, 3.7, 3.9, 2.6, "Чинники зараз", None, VIOLET),
         (4.75, 1.85, 4.05, 1.6, "Порадник", "загроза: Грінвальд\n«знизити уставку густини до 0,62, D₂ удвічі»\n[Застосувати]", EMERALD),
         (4.75, 3.7, 4.05, 2.6, "Автокерування", "дія, коли p > поріг\nпауза 0,3 с (плазма реагує)\nкожна дія — у журнал\n\n+ 10 повзунків уставок для ручного керування",
          CYAN))
for x, yy, w, h, head, body, c in panel:
    box(s, x, yy, w, h, CARD, line=c)
    paras = [(head, dict(size=15, bold=True, color=c, after=6))]
    if body:
        paras.append((body, dict(size=12, color=MUTED)))
    tb(s, x + 0.2, yy + 0.15, w - 0.4, h - 0.3, paras)
for i, (m, v, c) in enumerate((("Грінвальд", 78, ORANGE), ("радіація", 15, AMBER), ("q95", 5, INDIGO), ("інші", 2, DIM))):
    yy = 4.3 + i * 0.45
    tb(s, 0.8, yy, 1.2, 0.35, [(m, dict(size=11, color=MUTED))], anchor=MSO_ANCHOR.MIDDLE)
    box(s, 2.0, yy + 0.07, 2.2 * v / 100 + 0.05, 0.22, c, radius=0.5)
    tb(s, 2.1 + 2.2 * v / 100, yy, 0.7, 0.35, [(f"{v} %", dict(size=11, color=TEXT))], anchor=MSO_ANCHOR.MIDDLE)
rows = [["загроза", "дія порадника"], ["Грінвальд", "густина × 0,85; D₂ × 0,5"], ["Кінк", "Iₚ до q95 ≈ 3 (≥ 1 МА)"],
        ["Радіація", "неон → 0; NBI + 5 МВт"], ["Тройон", "NBI × 0,8"]]
table(s, 9.05, 1.85, 3.7, rows, [1.35, 2.35], row_h=0.62, sizes=[[12, 12]] + [[13, 12]] * 4)
tb(s, 9.05, 5.1, 3.7, 1.2, [("Повзунки: Iₚ, Bₜ, густина, D₂, неон, NBI, ICH, ECH, κ, δ", dict(size=12, color=MUTED))])
tb(s, 0.6, 6.6, 12, 0.4, [("Схема панелі «Керування» симулятора fusion-sim (JET, крок 5 мс); значення — ілюстрація", dict(size=11, color=DIM))])

# ───────────────────────── 14. замкнений контур ─────────────────────────
s = new_slide("ЗАМКНЕНИЙ КОНТУР", "З автокеруванням зривів удвічі менше", ORANGE, notes=(
    "≈45 с. Замкнений контур: ті самі 40 розрядів без керування і з ним. Модель тривоги бачить загрозу, ImpactNet називає "
    "механізм, порадник діє. Було 11 зривів, стало 5. Чесно: це 40 розрядів на безшумних каналах — наступний крок — 200 "
    "розрядів на тих самих сидах і облік ціни керування: на скільки впали струм і енергія."))
boxed_text(s, 0.6, 1.85, 5.6, 3.55, [("11 → 5", dict(size=80, bold=True, color=ORANGE, after=4)),
                                     ("зривів на 40 розрядах  (−55 %)", dict(size=18, color=TEXT, after=6)),
                                     ("модель v8 · поріг 0,25 · сид 3 · ті самі розряди без керування і з ним",
                                      dict(size=13, color=MUTED))], pad=0.4)
# коло: Прилади → Тривога ↓ Причина ← Дія ↑ Прилади
loop = ((6.65, 1.85, "Прилади", CYAN), (9.8, 1.85, "Тривога\n(модель)", EMERALD),
        (9.8, 3.7, "Причина\n(ImpactNet)", VIOLET), (6.65, 3.7, "Дія\n(порадник)", ORANGE))
for x, yy, t, c in loop:
    box(s, x, yy, 2.7, 1.35, CARD, line=c)
    tb(s, x, yy, 2.7, 1.35, [(t, dict(size=15, bold=True, color=c, align=PP_ALIGN.CENTER))], anchor=MSO_ANCHOR.MIDDLE)
arrow(s, 9.4, 2.37, w=0.4, h=0.3)
down_arrow(s, 11.0, 3.25, w=0.3, h=0.42)
box(s, 9.4, 4.22, 0.4, 0.3, DIM, shape=MSO_SHAPE.LEFT_ARROW)
box(s, 7.85, 3.25, 0.3, 0.42, DIM, shape=MSO_SHAPE.UP_ARROW)
tb(s, 6.65, 5.55, 6.1, 0.4, [("плазма реагує → прилади → … (кожні 5 мс)", dict(size=12, color=MUTED, align=PP_ALIGN.CENTER))])
chips(s, 0.6, 5.85, (("попередньо: 40 розрядів", AMBER), ("канали без шуму", AMBER), ("далі: 200 розрядів + ціна керування", MUTED)))

# ───────────────────────── 15. фінал ─────────────────────────
s = new_slide("ПІДСУМОК", "Коли, чому і що робити — у реальному часі", EMERALD, notes=(
    "≈40 с. Підсумок: модель тривоги без формул і без підказок — 0,924 проти 0,765 і 98 % зривів за 30 мс. ImpactNet "
    "пояснює причину — 49 з 49. Симулятор із порадником удвічі зменшує зриви. Далі: перевірка контуру на більшій вибірці "
    "і з шумом приладів, ціна керування, довчання на реальних даних. Код і результати — на GitHub. Дякуємо!"))
for i, (big, h, t, c) in enumerate((("0,924", "Коли", "модель тривоги без формул і підказок;\n98 % зривів за 30 мс", EMERALD),
                                    ("49 / 49", "Чому", "ImpactNet: частки 5 механізмів,\nта сама λ, що в симуляторі", VIOLET),
                                    ("11 → 5", "Що робити", "порадник і автокерування\nу замкненому контурі", ORANGE))):
    x = 0.6 + i * 3.0
    boxed_text(s, x, 1.85, 2.8, 2.35, [(big, dict(size=32, bold=True, color=c, after=4)), (h, dict(size=16, bold=True, after=4)),
                                       (t, dict(size=12, color=MUTED))])
tb(s, 0.6, 4.7, 8.6, 0.4, [("Далі", dict(size=16, bold=True, color=AMBER))])
tb(s, 0.6, 5.1, 8.6, 1.6, [("контур на 200+ розрядах і з шумом приладів · ціна керування (струм, енергія) · "
                            "довчання ImpactNet лише з моментів зривів на реальних даних", dict(size=13, color=MUTED))])
qr = qrcode.QRCode(border=1, box_size=10)
qr.add_data("https://" + REPO)
qr.make(fit=True)
qpath = OUT.parent / "repo_qr.png"
qr.make_image(fill_color="black", back_color="white").save(qpath)
box(s, 9.75, 1.85, 3.0, 3.0, "FFFFFF", radius=0.06)
s.shapes.add_picture(str(qpath), Inches(9.9), Inches(2.0), Inches(2.7), Inches(2.7))
tb(s, 9.3, 5.0, 3.9, 0.9, [("Код, моделі, звіти заліку", dict(size=13, bold=True, align=PP_ALIGN.CENTER, after=2)),
                           (REPO, dict(size=11, color=CYAN, align=PP_ALIGN.CENTER))])
tb(s, 0.6, 6.55, 12, 0.5, [("Дякуємо! Питання?", dict(size=20, bold=True, color=EMERALD))])

assert _n[0] == TOTAL, _n[0]
prs.save(OUT)
print("збережено", OUT)
