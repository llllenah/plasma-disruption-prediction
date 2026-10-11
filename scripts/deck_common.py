"""Спільні засоби для наукової презентації: біле тло, Arial, нумеровані рисунки й таблиці (ДСТУ 3008)."""
import sys
from pathlib import Path

import qrcode  # noqa: F401  (використовується у скриптах слайдів)
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "presentation/plasma_disruption.pptx")
FIGS = OUT.parent / "figs"
REPO = "github.com/llllenah/plasma-disruption-prediction"
FONT = "Arial"
BG, CARD, LINE = "FFFFFF", "F8FAFC", "E2E8F0"
TEXT, MUTED, DIM = "0F172A", "475569", "94A3B8"
IND, EME, ORA, VIO, AMB, RED, CYAN = "1D4ED8", "047857", "C2410C", "6D28D9", "B45309", "B91C1C", "0E7490"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 30
_n = [0]
_fig = [0]
_tab = [0]


def rgb(h):
    return RGBColor.from_string(h)


def tb(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    for i, (content, o) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(o.get("after", 0))
        for text, ro in (content if isinstance(content, list) else [(content, {})]):
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = FONT
            f.size = Pt(ro.get("size", o.get("size", 15)))
            f.bold = ro.get("bold", o.get("bold", False))
            f.italic = ro.get("italic", o.get("italic", False))
            f.color.rgb = rgb(ro.get("color", o.get("color", TEXT)))
    return s


def box(slide, x, y, w, h, fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04, lw=1.0):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def points(slide, x, y, w, items, size=15, gap=10, lead_color=TEXT):
    """Абзаци з виділеним терміном: [(термін, пояснення)]."""
    paras = [([(lead + " – " if lead else "", dict(bold=True, color=lead_color)), (text, dict(color=TEXT))],
              dict(size=size, after=gap)) for lead, text in items]
    return tb(slide, x, y, w, 0.5 * len(items) + 1, paras)


def number(slide, x, y, w, big, small, color, size=34):
    tb(slide, x, y, w, 1.3, [(big, dict(size=size, bold=True, color=color, after=2)), (small, dict(size=12, color=MUTED))])


def pic(slide, name, x, y, w=None, h=None, border=False):
    path = FIGS / name
    W, H = Image.open(path).size
    if w and not h:
        h = w * H / W
    elif h and not w:
        w = h * W / H
    if border:
        box(slide, x - 0.03, y - 0.03, w + 0.06, h + 0.06, None, LINE, shape=MSO_SHAPE.RECTANGLE)
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h


def fig_cap(slide, x, y, w, text):
    _fig[0] += 1
    tb(slide, x, y, w, 0.4, [(f"Рисунок {_fig[0]} – {text}", dict(size=11, color=MUTED))])


def tab_cap(slide, x, y, w, text):
    _tab[0] += 1
    tb(slide, x, y, w, 0.35, [(f"Таблиця {_tab[0]} – {text}", dict(size=11, color=MUTED))])


def note(slide, x, y, w, text):
    tb(slide, x, y, w, 0.4, [(text, dict(size=10, color=DIM))])


def arrow(slide, x, y, w=0.4, h=0.3, shape=MSO_SHAPE.RIGHT_ARROW):
    box(slide, x, y, w, h, DIM, None, shape=shape)


def table(slide, x, y, w, rows, col_w, row_h=0.42, sizes=None, fills=None, bold_col0=True):
    gs = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = gs.table
    tblPr = gs._element.graphic.graphicData.tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = rgb("F1F5F9" if i == 0 else (fills[i][j] if fills else "FFFFFF"))
            c.margin_left = c.margin_right = Inches(0.1)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            c.text_frame.word_wrap = True
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            r.font.name = FONT
            r.font.size = Pt(sizes[i][j] if sizes else (12 if i == 0 else 13))
            r.font.bold = i == 0 or (j == 0 and bold_col0)
            r.font.color.rgb = rgb(MUTED if i == 0 else TEXT)
    return gs


def new_slide(title, notes=""):
    _n[0] += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    if title:
        tb(s, 0.6, 0.45, 12.1, 0.7, [(title, dict(size=26, bold=True))])
    tb(s, 12.3, 7.05, 0.5, 0.3, [(str(_n[0]), dict(size=10, color=DIM, align=PP_ALIGN.RIGHT))])
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s
